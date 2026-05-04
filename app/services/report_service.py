import io
from datetime import date
from typing import List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from sqlalchemy.orm import Session
from sqlalchemy import func
from app.models import Student, Attendance
from app.schemas.student import StudentAttendanceStats


class ReportService:
    def generate_attendance_stats(
        self, 
        db: Session, 
        branch: Optional[str] = None,
        semester: Optional[int] = None,
        date_from: Optional[date] = None,
        date_to: Optional[date] = None,
        subject: Optional[str] = None
    ) -> List[StudentAttendanceStats]:
        """Generate attendance statistics for students."""
        
        # Base query
        query = db.query(Student)
        
        if branch:
            query = query.filter(Student.branch == branch)
        if semester:
            query = query.filter(Student.semester == semester)
        
        #students = query.all()
        students = query.order_by(func.lower(Student.full_name).asc()).all()
        stats = []
        
        for student in students:
            # Build attendance query
            att_query = db.query(Attendance).filter(Attendance.student_id == student.id)
            
            if date_from:
                att_query = att_query.filter(Attendance.date >= date_from)
            if date_to:
                att_query = att_query.filter(Attendance.date <= date_to)
            if subject:
                att_query = att_query.filter(Attendance.subject == subject)
            
            total = att_query.count()
            present = att_query.filter(Attendance.status == "present").count()
            absent = total - present
            
            percentage = (present / total * 100) if total > 0 else 0.0
            
            stats.append(StudentAttendanceStats(
                student_id=student.id,
                full_name=student.full_name,
                roll_number=student.roll_number,
                branch=student.branch,
                semester=student.semester,
                total_classes=total,
                present_days=present,
                absent_days=absent,
                attendance_percentage=round(percentage, 2)
            ))
        
        return stats
    
    def generate_excel_report(
        self,
        stats: List[StudentAttendanceStats],
        title: str = "Attendance Report"
    ) -> bytes:
        """Generate an Excel report from attendance stats."""
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance Report"
        
        # Styles
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="4A90D9", end_color="4A90D9", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Title
        ws.merge_cells('A1:G1')
        ws['A1'] = title
        ws['A1'].font = Font(bold=True, size=16)
        ws['A1'].alignment = Alignment(horizontal="center")
        
        # Headers
        headers = ["Roll Number", "Name", "Branch", "Semester", "Total Classes", "Present", "Attendance %"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Data
        for row, stat in enumerate(stats, 4):
            ws.cell(row=row, column=1, value=stat.roll_number).border = thin_border
            ws.cell(row=row, column=2, value=stat.full_name).border = thin_border
            ws.cell(row=row, column=3, value=stat.branch).border = thin_border
            ws.cell(row=row, column=4, value=stat.semester).border = thin_border
            ws.cell(row=row, column=5, value=stat.total_classes).border = thin_border
            ws.cell(row=row, column=6, value=stat.present_days).border = thin_border
            
            # Color code attendance percentage
            percentage_cell = ws.cell(row=row, column=7, value=f"{stat.attendance_percentage}%")
            percentage_cell.border = thin_border
            if stat.attendance_percentage >= 75:
                percentage_cell.fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
            elif stat.attendance_percentage >= 50:
                percentage_cell.fill = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
            else:
                percentage_cell.fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
        
        # Adjust column widths
        column_widths = [15, 25, 15, 12, 15, 10, 15]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    
    def generate_pdf_report(
        self,
        stats: List[StudentAttendanceStats],
        title: str = "Attendance Report"
    ) -> bytes:
        """Generate a PDF report from attendance stats."""
        
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        
        elements = []
        styles = getSampleStyleSheet()
        
        # Title
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            spaceAfter=30,
            alignment=1  # Center
        )
        elements.append(Paragraph(title, title_style))
        elements.append(Spacer(1, 0.25 * inch))
        
        # Table data
        data = [["Roll No.", "Name", "Branch", "Sem", "Total", "Present", "Absent", "Attendance %"]]
        
        for stat in stats:
            data.append([
                stat.roll_number,
                stat.full_name,
                stat.branch,
                str(stat.semester),
                str(stat.total_classes),
                str(stat.present_days),
                str(stat.absent_days),
                f"{stat.attendance_percentage}%"
            ])
        
        # Create table
        table = Table(data, repeatRows=1)
        
        # Table styling
        style = TableStyle([
            # Header
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4A90D9')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 11),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            
            # Body
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 10),
            ('ALIGN', (0, 1), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            
            # Grid
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            
            # Alternating row colors
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')]),
        ])
        
        # Color code attendance percentage column
        for i, stat in enumerate(stats, 1):
            if stat.attendance_percentage >= 75:
                style.add('BACKGROUND', (-1, i), (-1, i), colors.HexColor('#90EE90'))
            elif stat.attendance_percentage >= 50:
                style.add('BACKGROUND', (-1, i), (-1, i), colors.HexColor('#FFD700'))
            else:
                style.add('BACKGROUND', (-1, i), (-1, i), colors.HexColor('#FF6B6B'))
        
        table.setStyle(style)
        elements.append(table)
        
        # Build PDF
        doc.build(elements)
        output.seek(0)
        return output.getvalue()


# Global instance
report_service = ReportService()
